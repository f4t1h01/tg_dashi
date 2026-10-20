import datetime as dt
import io

import openpyxl
import pytest

from app.analytics import ABSENT, LATE, ON_TIME, WEEKEND_WORK, build_attendance, evaluate
from app.config import ReportSettings
from app.parser import ParseError, parse_workbook
from app.progress import MIN_DURATION, TICK, next_percent
from app.report import build_report, summary_text

MON = dt.datetime(2026, 10, 5)


def make_export(rows: list[tuple]) -> bytes:
    """Книга в формате выгрузки «Passages through perimeter»."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Passages through perimeter"])
    ws.append([])
    ws.append([None, None, "Creation date:", dt.datetime(2026, 10, 20, 20, 4)])
    ws.append([None, None, "Begin date:", MON])
    ws.append([None, None, "End date:", dt.datetime(2026, 10, 11, 23, 59)])
    ws.append([])
    ws.append(["Date", "Department", "Name", "Personnel Number", "Events"])
    ws.append([None, None, None, None, "arrival", None, "leaving"])
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ROWS = [
    # Пн: Иванов — продолжения дня без даты/ФИО, общежитие раньше КПП
    (MON, "Lab   ", "IVANOV", "001", "08:30", "Dorm3", "(no)", None),
    (None, None, None, None, "08:55", "KPP1 Turniket1", "12:00", "KPP1 Turniket2"),
    (None, None, None, None, "(no)", None, "18:30", "KPP 1 Exit"),
    # Пн: Петров опоздал и ушёл рано, первая отметка — выход
    (MON, "Lab", "PETROV", None, "(no)", None, "09:20", "Turniket 2"),
    (None, None, None, None, "09:25", "Turniket 1", "17:00", "Turniket 2"),
    # Сб: Иванов вышел на работу
    (dt.datetime(2026, 10, 10), "Lab", "IVANOV", "001", "10:00", "KPP2 Turniket1", "14:00", "KPP2 Turniket2"),
]


@pytest.fixture
def att():
    return build_attendance(parse_workbook(make_export(ROWS)), ReportSettings())


def test_parser_groups_continuation_rows():
    exp = parse_workbook(make_export(ROWS))
    assert exp.begin == dt.date(2026, 10, 5) and exp.end == dt.date(2026, 10, 11)
    ivanov = next(r for r in exp.records if r.name == "IVANOV" and r.date == MON.date())
    assert ivanov.department == "Lab"  # лишние пробелы убраны
    assert [e.time.strftime("%H:%M") for e in ivanov.events] == ["08:30", "08:55", "12:00", "18:30"]


def test_parser_rejects_foreign_file():
    wb = openpyxl.Workbook()
    wb.active.append(["что-то", "другое"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ParseError):
        parse_workbook(buf.getvalue())


def test_statuses_and_modes(att):
    by = {(r.date, r.employee.name): r for r in att.rows}
    ivanov = by[(MON.date(), "IVANOV")]
    gate_only = evaluate(ivanov, att.settings, mode=1)
    with_dorm = evaluate(ivanov, att.settings, mode=2)
    assert gate_only.passes.first_in == dt.time(8, 55) and gate_only.status == ON_TIME
    assert with_dorm.passes.first_in == dt.time(8, 30)
    assert gate_only.passes.last_out == dt.time(18, 30) and gate_only.early_min == 0

    petrov = evaluate(by[(MON.date(), "PETROV")], att.settings)
    assert petrov.passes.first_in == dt.time(9, 20)  # первая отметка — выход
    assert petrov.status == LATE and petrov.late_min == 20 and petrov.early_min == 60

    assert evaluate(by[(dt.date(2026, 10, 6), "PETROV")], att.settings).status == ABSENT
    assert evaluate(by[(dt.date(2026, 10, 10), "IVANOV")], att.settings).status == WEEKEND_WORK
    assert (dt.date(2026, 10, 11), "IVANOV") not in by  # выходной без проходов — строки нет


def test_late_grace(att):
    s = ReportSettings(late_grace_min=30)
    row = next(r for r in att.rows if r.employee.name == "PETROV" and r.date == MON.date())
    assert evaluate(row, s).status == ON_TIME


def test_report_builds(att):
    data = build_report(att)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Сводка", "За месяц", "За день", "_calc"]
    assert wb["_calc"].sheet_state == "hidden"
    cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True)["За день"]
    statuses = [cached.cell(r, 6).value for r in range(5, 5 + len(att.rows))]
    assert statuses.count(LATE) == 1 and statuses.count(WEEKEND_WORK) == 1
    for lang in ("ru", "uz", "en"):
        assert "PETROV" in summary_text(att, lang)


@pytest.mark.parametrize("work", [0.1, 3.0, 20.0])
def test_progress_min_duration_and_slowdown(work):
    p, elapsed, seen_50 = 0.0, 0.0, None
    while p < 100:
        p = next_percent(p, elapsed >= work, elapsed)
        if p >= 50 and seen_50 is None:
            seen_50 = elapsed
        if elapsed < work:
            assert p < 96  # пока работа идёт, 100% не показываем
        elapsed += TICK
    assert elapsed >= MIN_DURATION
    assert elapsed >= work
